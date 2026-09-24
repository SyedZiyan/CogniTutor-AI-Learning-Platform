import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MATERIALS_DIR = os.path.join(BASE_DIR, "sample_materials")
os.makedirs(MATERIALS_DIR, exist_ok=True)

# 1. Text File: Deep Learning Fundamentals
txt_content = """# Deep Learning & Neural Networks Fundamentals

## Chapter 1: Introduction to Artificial Neural Networks
Artificial Neural Networks (ANNs) are computational models inspired by the biological neural networks of the human brain. An ANN consists of interconnected nodes called artificial neurons or perceptrons organized into layers: the input layer, one or more hidden layers, and the output layer.

Each neuron computes a weighted sum of its inputs, adds a bias term, and passes the result through a non-linear activation function:
z = sum(w_i * x_i) + b
a = sigma(z)

Common activation functions include:
1. Sigmoid: Maps values between 0 and 1. Useful for binary classification, but suffers from vanishing gradients.
2. ReLU (Rectified Linear Unit): f(x) = max(0, x). Solves vanishing gradient for positive values and accelerates gradient descent convergence.
3. Leaky ReLU: f(x) = max(0.01x, x). Prevents dying ReLU problem by providing a small non-zero slope for negative inputs.
4. Softmax: Converts raw logits into normalized probability distribution over multiple classes.

## Chapter 2: Forward Propagation and Loss Functions
Forward propagation is the process where input data traverses forward through the layers to compute the network's prediction. Once the prediction is made, a loss function measures how far the prediction is from the actual ground truth.
- Mean Squared Error (MSE): Typically used for continuous regression tasks.
- Binary Cross-Entropy: Used for binary classification tasks.
- Categorical Cross-Entropy: Used for multi-class classification tasks.

## Chapter 3: Backpropagation and Gradient Descent
Backpropagation (backward propagation of errors) is the fundamental algorithm used to train neural networks. It applies the chain rule of calculus to compute the partial derivative of the loss function with respect to every weight and bias in the network:
dL/dw_ij = (dL/da_j) * (da_j/dz_j) * (dz_j/dw_ij)

Once gradients are computed, an optimization algorithm updates the parameters:
w_new = w_old - learning_rate * dL/dw

Key optimizers include:
- Stochastic Gradient Descent (SGD) with Momentum
- RMSProp: Scales the learning rate based on running average of squared gradients
- Adam (Adaptive Moment Estimation): Combines momentum and RMSProp for robust convergence.

## Chapter 4: Overfitting, Underfitting, and Regularization
Overfitting occurs when a neural network learns noise, idiosyncrasies, and specific patterns in the training data so closely that it performs exceptionally well on training data (low training error) but fails to generalize to unseen validation or test data (high test error).
Conversely, underfitting happens when the model is too simple to capture underlying patterns, resulting in high error on both training and test data.

Techniques to prevent overfitting:
1. Dropout: Randomly deactivates a fraction p of neurons during each training step, preventing co-adaptation of features.
2. L1 (Lasso) and L2 (Ridge / Weight Decay) Regularization: Adds a penalty term proportional to weights to the loss function.
3. Early Stopping: Halts training when validation loss stops improving.
4. Data Augmentation: Synthetically expands the dataset by rotating, cropping, or flipping inputs.
"""

with open(os.path.join(MATERIALS_DIR, "Deep_Learning_Fundamentals.txt"), "w", encoding="utf-8") as f:
    f.write(txt_content)
print("Created Deep_Learning_Fundamentals.txt")

# 2. PDF File: Convolutional Neural Networks
pdf_path = os.path.join(MATERIALS_DIR, "Convolutional_Neural_Networks.pdf")
doc = SimpleDocTemplate(pdf_path, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
styles = getSampleStyleSheet()

title_style = ParagraphStyle(
    'DocTitle',
    parent=styles['Heading1'],
    fontSize=20,
    leading=24,
    textColor='#1E3A8A'
)
h2_style = ParagraphStyle(
    'Heading2',
    parent=styles['Heading2'],
    fontSize=13,
    leading=17,
    textColor='#2563EB',
    spaceBefore=10,
    spaceAfter=4
)
body_style = ParagraphStyle(
    'Body',
    parent=styles['Normal'],
    fontSize=10,
    leading=14,
    textColor='#1F2937'
)

story = []
story.append(Paragraph("Chapter 4: Convolutional Neural Networks (CNNs)", title_style))
story.append(Spacer(1, 8))
story.append(Paragraph("A Comprehensive Guide to Spatial Feature Extraction and Computer Vision", body_style))
story.append(Spacer(1, 12))

story.append(Paragraph("1. Motivation for CNNs in Computer Vision", h2_style))
story.append(Paragraph(
    "Standard Fully Connected (Dense) networks struggle with image data because flattening a 2D image (e.g., 256x256x3) creates an enormous number of parameters (nearly 200,000 weights per hidden neuron), leading to severe overfitting and computational bottlenecks. Furthermore, dense networks discard spatial 2D relationships between neighboring pixels. CNNs solve these challenges through three foundational principles: local receptive fields, shared weights (parameter sharing), and spatial subsampling.",
    body_style
))

story.append(Paragraph("2. The Convolution Operation and Filters", h2_style))
story.append(Paragraph(
    "The core building block of a CNN is the convolutional layer. A set of learnable kernels (filters), typically 3x3 or 5x5, slides across the input tensor. At each location, element-wise multiplication is performed followed by a summation, producing a 2D activation map (feature map). Key parameters include stride (step size of filter movement) and padding (adding zeros around borders to maintain spatial dimensions, known as 'same' padding).",
    body_style
))

story.append(Paragraph("3. Pooling Layers (Downsampling)", h2_style))
story.append(Paragraph(
    "Pooling layers reduce the spatial dimensions (height and width) of feature maps while retaining essential information. Max Pooling takes the maximum value in a window (e.g. 2x2 with stride 2), providing translation invariance and drastically reducing parameter count and computational complexity. Average Pooling calculates the mean value across the window.",
    body_style
))

story.append(Paragraph("4. Landmark Architectures: LeNet to ResNet", h2_style))
story.append(Paragraph(
    "Over the years, landmark CNN architectures have driven progress in computer vision: LeNet-5 (1998, digit recognition), AlexNet (2012, ImageNet breakthrough using ReLU and Dropout), VGGNet (2014, deep stacks of small 3x3 convolutions), and ResNet (2015, residual skip connections solving the degradation and vanishing gradient problem in ultra-deep networks of 50-152 layers).",
    body_style
))

story.append(Paragraph("5. Advantages of CNNs", h2_style))
story.append(Paragraph(
    "The primary advantages of CNNs mentioned in this chapter are: (a) Parameter sharing reduces model complexity compared to dense layers; (b) Translation invariance allows detecting objects regardless of position; (c) Hierarchical feature learning automatically extracts edges in early layers, textures in middle layers, and complex semantic objects in deep layers without manual feature engineering.",
    body_style
))

doc.build(story)
print("Created Convolutional_Neural_Networks.pdf")

# 3. DOCX File: Recurrent Neural Networks and LSTMs
docx_path = os.path.join(MATERIALS_DIR, "Recurrent_Neural_Networks_and_LSTMs.docx")
doc_docx = Document()
doc_docx.add_heading("Recurrent Neural Networks, LSTMs, and Sequential Modeling", level=0)

doc_docx.add_heading("1. Sequential Data and Vanilla RNNs", level=1)
doc_docx.add_paragraph(
    "Sequential data such as text, speech, audio, and time-series involves temporal dependencies where the current element depends on preceding context. Vanilla Recurrent Neural Networks (RNNs) address this by maintaining a hidden state vector h_t that functions as working memory. At time step t, the hidden state is updated as:\n"
    "h_t = tanh(W_hh * h_{t-1} + W_xh * x_t + b_h)\n"
    "y_t = softmax(W_hy * h_t + b_y)"
)

doc_docx.add_heading("2. The Vanishing and Exploding Gradient Problem", level=1)
doc_docx.add_paragraph(
    "During training via Backpropagation Through Time (BPTT), gradients are repeatedly multiplied across sequence time steps. If the eigenvalues of the recurrent weight matrix W_hh are less than 1, gradients shrink exponentially with sequence length, causing the vanishing gradient problem. The network fails to learn long-range dependencies and forgets early inputs. Conversely, if eigenvalues exceed 1, gradients explode causing numerical overflow (mitigated by gradient clipping)."
)

doc_docx.add_heading("3. Long Short-Term Memory (LSTM) Networks", level=1)
doc_docx.add_paragraph(
    "Hochreiter and Schmidhuber (1997) introduced the Long Short-Term Memory (LSTM) architecture to overcome vanishing gradients. LSTMs introduce an explicit cell state C_t that acts as a conveyor belt passing information unchanged across time steps, regulated by three specialized gates:\n"
    "1. Forget Gate (f_t = sigmoid(W_f * [h_{t-1}, x_t] + b_f)): Decides what proportion of information to discard from the previous cell state.\n"
    "2. Input Gate (i_t = sigmoid(W_i * [h_{t-1}, x_t] + b_i)) and Candidate State (~C_t = tanh(W_c * [h_{t-1}, x_t] + b_c)): Decides which new values to write into memory.\n"
    "3. Cell State Update: C_t = f_t * C_{t-1} + i_t * ~C_t\n"
    "4. Output Gate (o_t = sigmoid(W_o * [h_{t-1}, x_t] + b_o)) and Hidden State: h_t = o_t * tanh(C_t)."
)

doc_docx.add_heading("4. Gated Recurrent Units (GRU)", level=1)
doc_docx.add_paragraph(
    "Cho et al. (2014) proposed GRUs as a streamlined alternative to LSTMs. GRUs combine the cell state and hidden state, and utilize only two gates: the Reset Gate (r_t) and the Update Gate (z_t). GRUs have fewer parameters, train faster, and frequently achieve comparable performance on small to medium datasets."
)

doc_docx.save(docx_path)
print("Created Recurrent_Neural_Networks_and_LSTMs.docx")

# 4. PPTX File: Transformers and Attention Mechanisms
pptx_path = os.path.join(MATERIALS_DIR, "Transformers_and_Attention_Mechanisms.pptx")
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

# Slide 1: Title
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
slide1.shapes.title.text = "Transformers & Modern Attention Mechanisms"
slide1.placeholders[1].text = "Attention Is All You Need | Multi-Head Self-Attention & LLM Foundations"

# Slide 2: Limitations of RNNs
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "Slide 2: Limitations of Recurrent Architectures"
content2 = slide2.placeholders[1].text_frame
content2.text = "Why move beyond RNNs and LSTMs?"
p1 = content2.add_paragraph()
p1.text = "• Sequential bottleneck: RNNs compute step-by-step (t-1 -> t), preventing parallel GPU training."
p2 = content2.add_paragraph()
p2.text = "• Information bottleneck: Compressing an entire sentence into a fixed-length vector loses fine details."
p3 = content2.add_paragraph()
p3.text = "• Long path length: Information must traverse O(N) sequential steps between tokens."

# Slide 3: Scaled Dot-Product Attention
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "Slide 3: Scaled Dot-Product Attention"
content3 = slide3.placeholders[1].text_frame
content3.text = "The Mathematical Core of Transformers:"
p1 = content3.add_paragraph()
p1.text = "• Every token is projected into Query (Q), Key (K), and Value (V) matrices."
p2 = content3.add_paragraph()
p2.text = "• Formula: Attention(Q, K, V) = softmax( (Q * K^T) / sqrt(d_k) ) * V"
p3 = content3.add_paragraph()
p3.text = "• The scaling factor sqrt(d_k) prevents large dot products from pushing softmax into vanishing gradient regions."

# Slide 4: Multi-Head Attention & Positional Encoding
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
slide4.shapes.title.text = "Slide 4: Multi-Head Attention & Positional Encoding"
content4 = slide4.placeholders[1].text_frame
content4.text = "Capturing Diverse Relationships:"
p1 = content4.add_paragraph()
p1.text = "• Multi-Head Attention projects Q, K, V into h different subspaces, enabling the model to jointly attend to syntactic, semantic, and positional aspects."
p2 = content4.add_paragraph()
p2.text = "• Positional Encoding (sinusoidal or learned) injects order information since attention is permutation invariant."
p3 = content4.add_paragraph()
p3.text = "• Modern LLMs (BERT = Encoder only, GPT = Decoder only, T5 = Encoder-Decoder)."

prs.save(pptx_path)
print("Created Transformers_and_Attention_Mechanisms.pptx")
