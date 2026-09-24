import os
import re
import uuid
from typing import List, Dict, Any
import pypdf
from pptx import Presentation
from docx import Document

class DocumentChunk:
    def __init__(self, chunk_id: str, doc_id: str, doc_name: str, text: str, page_or_slide: str, section: str, topic: str):
        self.chunk_id = chunk_id
        self.doc_id = doc_id
        self.doc_name = doc_name
        self.text = text.strip()
        self.page_or_slide = page_or_slide
        self.section = section
        self.topic = topic

    def to_dict(self) -> Dict[str, Any]:
        return {
            "chunk_id": self.chunk_id,
            "doc_id": self.doc_id,
            "doc_name": self.doc_name,
            "text": self.text,
            "page_or_slide": self.page_or_slide,
            "section": self.section,
            "topic": self.topic
        }

class DocumentProcessor:
    KNOWN_TOPICS = {
        "Neural Networks": ["neural network", "perceptron", "layer", "neuron", "activation function", "sigmoid", "relu", "softmax"],
        "Forward & Loss Functions": ["forward propagation", "loss function", "mse", "cross-entropy", "mean squared error"],
        "Backpropagation & Optimization": ["backpropagation", "gradient descent", "chain rule", "sgd", "rmsprop", "adam", "learning rate", "optimizer"],
        "Overfitting & Regularization": ["overfitting", "underfitting", "regularization", "dropout", "l1", "l2", "weight decay", "early stopping", "generalization"],
        "Convolutional Neural Networks (CNN)": ["cnn", "convolution", "kernel", "filter", "feature map", "stride", "padding", "pooling", "lenet", "alexnet", "vgg", "resnet"],
        "Recurrent Neural Networks (RNN)": ["rnn", "recurrent", "sequential", "time step", "hidden state", "bptt", "backpropagation through time"],
        "Vanishing Gradient Problem": ["vanishing gradient", "exploding gradient", "eigenvalues", "gradient clipping", "long-range dependencies"],
        "Long Short-Term Memory (LSTM)": ["lstm", "long short-term memory", "cell state", "forget gate", "input gate", "output gate", "candidate state"],
        "Gated Recurrent Units (GRU)": ["gru", "gated recurrent unit", "reset gate", "update gate"],
        "Attention & Transformers": ["transformer", "attention", "self-attention", "multi-head attention", "query", "key", "value", "positional encoding", "bert", "gpt"]
    }

    @staticmethod
    def detect_topic(text: str) -> str:
        text_lower = text.lower()
        topic_scores = {}
        for topic, keywords in DocumentProcessor.KNOWN_TOPICS.items():
            score = sum(1 for kw in keywords if re.search(r'\b' + re.escape(kw) + r'\b', text_lower))
            if score > 0:
                topic_scores[topic] = score
        if topic_scores:
            return max(topic_scores.items(), key=lambda x: x[1])[0]
        # Fallback to general topic
        return "Deep Learning Fundamentals"

    @classmethod
    def parse_pdf(cls, file_path: str, doc_id: str, doc_name: str) -> List[DocumentChunk]:
        chunks = []
        reader = pypdf.PdfReader(file_path)
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if not page_text.strip():
                continue
            page_label = f"Page {page_idx + 1}"
            section_title = f"{doc_name} - {page_label}"
            lines = [l.strip() for l in page_text.splitlines() if l.strip()]
            if lines and len(lines[0]) < 60:
                section_title = lines[0]

            page_chunks = cls.split_into_chunks(page_text, chunk_size=600, overlap=100)
            for i, chunk_text in enumerate(page_chunks):
                topic = cls.detect_topic(chunk_text)
                chunk = DocumentChunk(
                    chunk_id=f"{doc_id}_p{page_idx+1}_c{i+1}",
                    doc_id=doc_id,
                    doc_name=doc_name,
                    text=chunk_text,
                    page_or_slide=page_label,
                    section=section_title,
                    topic=topic
                )
                chunks.append(chunk)
        return chunks

    @classmethod
    def parse_pptx(cls, file_path: str, doc_id: str, doc_name: str) -> List[DocumentChunk]:
        chunks = []
        prs = Presentation(file_path)
        for slide_idx, slide in enumerate(prs.slides):
            slide_label = f"Slide {slide_idx + 1}"
            slide_title = f"{doc_name} - {slide_label}"
            text_parts = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    if shape == slide.shapes.title:
                        slide_title = shape.text_frame.text.strip()
                    else:
                        text_parts.append(shape.text_frame.text.strip())

            full_text = f"{slide_title}\n" + "\n".join(text_parts)
            if not full_text.strip():
                continue

            slide_chunks = cls.split_into_chunks(full_text, chunk_size=550, overlap=80)
            for i, chunk_text in enumerate(slide_chunks):
                topic = cls.detect_topic(chunk_text)
                chunk = DocumentChunk(
                    chunk_id=f"{doc_id}_s{slide_idx+1}_c{i+1}",
                    doc_id=doc_id,
                    doc_name=doc_name,
                    text=chunk_text,
                    page_or_slide=slide_label,
                    section=slide_title,
                    topic=topic
                )
                chunks.append(chunk)
        return chunks

    @classmethod
    def parse_docx(cls, file_path: str, doc_id: str, doc_name: str) -> List[DocumentChunk]:
        chunks = []
        doc = Document(file_path)
        current_section = doc_name
        current_text_blocks = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                if current_text_blocks:
                    section_text = "\n".join(current_text_blocks)
                    sec_chunks = cls.split_into_chunks(section_text, chunk_size=600, overlap=100)
                    for i, ct in enumerate(sec_chunks):
                        topic = cls.detect_topic(ct)
                        chunks.append(DocumentChunk(
                            chunk_id=f"{doc_id}_sec{len(chunks)+1}_c{i+1}",
                            doc_id=doc_id,
                            doc_name=doc_name,
                            text=ct,
                            page_or_slide=f"Section: {current_section[:25]}",
                            section=current_section,
                            topic=topic
                        ))
                    current_text_blocks = []
                current_section = text
            else:
                current_text_blocks.append(text)

        if current_text_blocks:
            section_text = "\n".join(current_text_blocks)
            sec_chunks = cls.split_into_chunks(section_text, chunk_size=600, overlap=100)
            for i, ct in enumerate(sec_chunks):
                topic = cls.detect_topic(ct)
                chunks.append(DocumentChunk(
                    chunk_id=f"{doc_id}_sec{len(chunks)+1}_c{i+1}",
                    doc_id=doc_id,
                    doc_name=doc_name,
                    text=ct,
                    page_or_slide=f"Section: {current_section[:25]}",
                    section=current_section,
                    topic=topic
                ))

        return chunks

    @classmethod
    def parse_txt(cls, file_path: str, doc_id: str, doc_name: str) -> List[DocumentChunk]:
        chunks = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        sections = re.split(r'\n(?=#{1,3}\s+)', content)
        for sec_idx, sec in enumerate(sections):
            sec = sec.strip()
            if not sec:
                continue
            lines = sec.splitlines()
            header_match = re.match(r'^#{1,3}\s+(.+)', lines[0])
            section_title = header_match.group(1).strip() if header_match else f"Chapter {sec_idx+1}"
            body_text = "\n".join(lines[1:]) if header_match else sec

            sec_chunks = cls.split_into_chunks(body_text, chunk_size=600, overlap=100)
            for i, ct in enumerate(sec_chunks):
                topic = cls.detect_topic(f"{section_title} {ct}")
                chunk = DocumentChunk(
                    chunk_id=f"{doc_id}_ch{sec_idx+1}_c{i+1}",
                    doc_id=doc_id,
                    doc_name=doc_name,
                    text=f"[{section_title}]\n{ct}",
                    page_or_slide=f"Chapter {sec_idx+1}",
                    section=section_title,
                    topic=topic
                )
                chunks.append(chunk)
        return chunks

    @classmethod
    def split_into_chunks(cls, text: str, chunk_size: int = 600, overlap: int = 100) -> List[str]:
        text = re.sub(r'\s+', ' ', text).strip()
        if len(text) <= chunk_size:
            return [text] if text else []

        sentences = re.split(r'(?<=[.!?])\s+', text)
        chunks = []
        current_chunk = []
        current_len = 0

        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            if current_len + len(sentence) > chunk_size and current_chunk:
                chunks.append(" ".join(current_chunk))
                # retain last few sentences for overlap
                overlap_chunk = []
                overlap_len = 0
                for s in reversed(current_chunk):
                    if overlap_len + len(s) <= overlap:
                        overlap_chunk.insert(0, s)
                        overlap_len += len(s)
                    else:
                        break
                current_chunk = overlap_chunk + [sentence]
                current_len = overlap_len + len(sentence)
            else:
                current_chunk.append(sentence)
                current_len += len(sentence)

        if current_chunk:
            chunks.append(" ".join(current_chunk))

        return chunks

    @classmethod
    def process_file(cls, file_path: str, doc_name: str = None) -> List[DocumentChunk]:
        if not doc_name:
            doc_name = os.path.basename(file_path)
        doc_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, doc_name))[:8]
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return cls.parse_pdf(file_path, doc_id, doc_name)
        elif ext in [".pptx", ".ppt"]:
            return cls.parse_pptx(file_path, doc_id, doc_name)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path, doc_id, doc_name)
        elif ext in [".txt", ".md", ".csv"]:
            return cls.parse_txt(file_path, doc_id, doc_name)
        else:
            # Fallback to plain text reader
            return cls.parse_txt(file_path, doc_id, doc_name)
